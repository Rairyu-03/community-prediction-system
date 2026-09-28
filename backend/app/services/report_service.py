from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
from math import log10
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..defaults import AGE_GROUPS
from ..models import Community, PredictionResult, PredictionRun, Project


BLUE = "1976A3"
LIGHT_BLUE = "EAF5F8"
CATEGORIES = ["助餐设施", "养老服务设施", "托育服务设施", "停车充电设施", "公共活动场地"]
REPORT_FORMAT_VERSION = "bar-v2"


def _font(size: int):
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def _shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_text(cell, value: object, *, bold: bool = False, color: str | None = None, size: float = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    normal = document.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(9)
    for style_name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10)):
        style = document.styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
    header = section.header.paragraphs[0]
    header.text = "完整社区设施需求预测与动态匹配"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(120, 140, 150)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("预测结果仅作为规划决策辅助，最终方案应结合现场踏勘与主管部门审查。")
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor(130, 140, 145)


def _year_chunks(rows: list[PredictionResult], size: int = 5) -> list[list[PredictionResult]]:
    """Keep every forecast year while limiting each table to a readable width."""
    return [rows[index:index + size] for index in range(0, len(rows), size)]


def _set_column_widths(table, widths: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)


def _repeat_table_header(table) -> None:
    properties = table.rows[0]._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _format_value(label: str, value: object) -> str:
    if value in (None, ""):
        return "—"
    number = float(value)
    suffix = "%" if "占比" in label else "㎡" if "面积" in label else "人" if "人数" in label or "容纳量" in label else "个"
    rendered = f"{number:,.2f}".rstrip("0").rstrip(".")
    return f"{rendered}{suffix}"


def _population_chart_series(rows: list[PredictionResult]) -> list[dict[str, object]]:
    """Build the single source of truth for every population bar series in a report."""
    definitions = [
        ("总人口", "人", lambda row: row.data.get("totalPopulation", 0)),
        ("总户数", "户", lambda row: row.data.get("households", 0)),
        *[(group, "人", lambda row, group=group: row.data.get("ageGroups", {}).get(group, 0)) for group in AGE_GROUPS],
    ]
    return [
        {"name": name, "unit": unit, "values": [float(value_of(row) or 0) for row in rows]}
        for name, unit, value_of in definitions
    ]


def _population_chart(rows: list[PredictionResult], community: Community) -> BytesIO:
    width, height = 1600, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 32), f"{community.code}-{community.name} 人口规模与年龄结构年度对比", fill="#164C67", font=_font(36))
    draw.text((80, 84), "总人口、总户数及全部年龄段分组柱状图（对数尺度仅调整柱高）", fill="#607F8E", font=_font(20))
    years = [row.year for row in rows]
    series = _population_chart_series(rows)
    values = [value for item in series for value in item["values"]]
    colors = ["#078583", "#276CC3", "#8B5CF6", "#CA641B", "#D13769", "#24824C", "#AC8020", "#657591", "#965038", "#CC499B"]
    left, top, right, bottom = 130, 150, width - 80, 640
    maximum = max(values, default=1)
    transformed_max = max(1.0, log10(1 + maximum)) * 1.04
    transform = lambda value: log10(1 + max(0.0, value))
    inverse = lambda value: 10 ** value - 1
    slot_width = (right - left) / max(len(years), 1)
    group_width = min(slot_width * 0.78, 235)
    gap = 3
    bar_width = max(5, (group_width - gap * (len(series) - 1)) / len(series))
    for index in range(len(years)):
        draw.rounded_rectangle((left + index * slot_width + slot_width * 0.05, top, left + (index + 1) * slot_width - slot_width * 0.05, bottom), radius=12, fill="#F8FBFC")
    for index in range(5):
        y = top + (bottom - top) * index / 4
        value = inverse(transformed_max * (4 - index) / 4)
        draw.line((left, y, right, y), fill="#D9E7EC", width=1)
        draw.text((20, y - 12), f"{value:,.0f}", fill="#587887", font=_font(19))
    draw.line((left, top, left, bottom), fill="#7699AA", width=2)
    draw.line((left, bottom, right, bottom), fill="#7699AA", width=2)
    for index, year in enumerate(years):
        center = left + (index + 0.5) * slot_width
        draw.text((center - 28, bottom + 18), str(year), fill="#587887", font=_font(19))
    for series_index, item in enumerate(series):
        color = colors[series_index % len(colors)]
        for year_index, value in enumerate(item["values"]):
            group_left = left + year_index * slot_width + (slot_width - group_width) / 2
            x1 = group_left + series_index * (bar_width + gap)
            x2 = x1 + bar_width
            y = bottom - transform(value) / transformed_max * (bottom - top)
            draw.rounded_rectangle((x1, y, x2, bottom), radius=min(6, int(bar_width / 2)), fill=color)
    legend_top = 705
    for index, item in enumerate(series):
        column, row = index % 5, index // 5
        x, y = 85 + column * 300, legend_top + row * 62
        color = colors[index % len(colors)]
        draw.rounded_rectangle((x, y + 3, x + 24, y + 23), radius=5, fill=color)
        draw.text((x + 50, y), f"{item['name']}（{item['unit']}）", fill="#244F62", font=_font(21))
    draw.text((85, height - 54), "说明：纵轴显示实际数量的对数刻度，图例颜色与柱体一一对应。", fill="#6C8793", font=_font(18))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _add_population_section(document: Document, rows: list[PredictionResult], community: Community) -> None:
    document.add_heading("1 人口预测结果", level=2)
    chart = document.add_picture(_population_chart(rows, community), width=Inches(6.8))
    chart._inline.docPr.set("descr", "人口年度分组柱状图：包含总人口、总户数及8个年龄段，共10组预测柱体")
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph("图 1  人口规模与年龄结构年度分组柱状图（8个年龄段、总人口、总户数；纵轴采用对数尺度，数值均为实际预测值）")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    age_groups = list(AGE_GROUPS)
    for selected in _year_chunks(rows):
        if len(rows) > 5:
            document.add_paragraph(f"人口结构（{selected[0].year}-{selected[-1].year}年）")
        table = document.add_table(rows=1, cols=1 + len(selected))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        _set_column_widths(table, [3.0, *([2.7] * len(selected))])
        _repeat_table_header(table)
        for index, header in enumerate(["年龄段", *[f"{row.year}年" for row in selected]]):
            _set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
            _shade(table.rows[0].cells[index], BLUE)
        for group in [*age_groups, "总人口", "总户数"]:
            cells = table.add_row().cells
            _set_cell_text(cells[0], group, bold=group in {"总人口", "总户数"})
            for index, row in enumerate(selected, start=1):
                value = row.data.get("totalPopulation", 0) if group == "总人口" else row.data.get("households", 0) if group == "总户数" else row.data.get("ageGroups", {}).get(group, 0)
                _set_cell_text(cells[index], f"{int(value):,}")


def _add_facility_table(document: Document, category: str, rows: list[PredictionResult]) -> None:
    document.add_heading(f"2.{CATEGORIES.index(category) + 1} {category}预测结果", level=3)
    metric_names = [name for name in rows[0].data.get("metrics", {}) if name.startswith(f"{category}/")]
    references = rows[-1].data.get("references", {})
    for selected in _year_chunks(rows):
        if len(rows) > 5:
            document.add_paragraph(f"{category}（{selected[0].year}-{selected[-1].year}年）")
        table = document.add_table(rows=1, cols=1 + len(selected) + 2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_column_widths(table, [3.4, *([1.65] * len(selected)), 2.8, 3.2])
        _repeat_table_header(table)
        headers = ["参数项", *[f"{row.year}年" for row in selected], "参考标准", "参考来源"]
        for cell, header in zip(table.rows[0].cells, headers):
            _set_cell_text(cell, header, bold=True, color="FFFFFF")
            _shade(cell, BLUE)
        for metric in metric_names:
            label = metric.split("/", 1)[1]
            cells = table.add_row().cells
            _set_cell_text(cells[0], label, size=7.1)
            for index, row in enumerate(selected, start=1):
                _set_cell_text(cells[index], _format_value(label, row.data.get("metrics", {}).get(metric)), size=7.1)
            detail = references.get(metric, {})
            _set_cell_text(cells[-2], detail.get("standard", "按用户参数测算"), size=6.5)
            _set_cell_text(cells[-1], detail.get("source", "用户录入，来源未填写"), size=6.5)
        for row_index, row in enumerate(table.rows):
            if row_index % 2 == 0 and row_index:
                for cell in row.cells:
                    _shade(cell, LIGHT_BLUE)


def _add_community_report(document: Document, index: int, community: Community, facility_rows: list[PredictionResult], population_rows: list[PredictionResult]) -> None:
    document.add_heading(f"报告{index}：{community.code}-{community.name}", level=1)
    info = document.add_table(rows=2, cols=4)
    info.style = "Table Grid"
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = ["社区编号", community.code, "社区名称", community.name, "参数参考地区", community.region, "预测时间范围", f"{facility_rows[0].year}-{facility_rows[-1].year}年"]
    for cell, value in zip([cell for row in info.rows for cell in row.cells], values):
        _set_cell_text(cell, value, bold=value in {"社区编号", "社区名称", "参数参考地区", "预测时间范围"})
    _add_population_section(document, population_rows, community)
    document.add_heading("2 设施需求预测结果", level=2)
    for category in CATEGORIES:
        _add_facility_table(document, category, facility_rows)
    document.add_heading("3 选址建议", level=2)
    document.add_paragraph(
        "分析过程：采用报告末期目标年结果，以设施面积需求减去社区现状供给得到缺口；正值进入新增选址分析，负值表示现有供给富余。随后逐条比较已导入闲置空间的建筑面积和用地面积。"
        "满足全部面积门槛时选择面积余量较小的候选空间，均不满足时给出覆盖率最高的空间及剩余缺口，并明确需要组合选址或跨社区统筹。"
    )
    suggestions = facility_rows[-1].data.get("siteSuggestions", {})
    for category, suggestion in suggestions.items():
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{category}：").bold = True
        paragraph.add_run(str(suggestion))
    document.add_heading("4 备注说明", level=2)
    document.add_paragraph("缺口类指标为正值时表示供给不足，负值表示供给富余，0表示供需平衡。坐标目前用于标识候选空间，并未用于路网距离计算；落地前仍需进行权属、交通、安全、服务半径和规划合规性复核。")


def generate_facility_report(db: Session, project: Project, run: PredictionRun, mode: str, community_id: int | None = None) -> Path:
    results = list(db.scalars(select(PredictionResult).where(PredictionResult.run_id == run.id).order_by(PredictionResult.community_id, PredictionResult.year)).all())
    if not results:
        raise ValueError("设施预测结果为空，无法生成报告")
    grouped: dict[int, list[PredictionResult]] = defaultdict(list)
    for result in results:
        grouped[result.community_id].append(result)
    if mode == "community":
        selected_id = community_id or next(iter(grouped))
        if selected_id not in grouped:
            raise ValueError("所选社区不在本次设施预测结果中")
        grouped = {selected_id: grouped[selected_id]}

    population_run_id = run.parameters_snapshot.get("_populationRunId")
    population_rows = list(db.scalars(select(PredictionResult).where(PredictionResult.run_id == population_run_id).order_by(PredictionResult.community_id, PredictionResult.year)).all())
    population_grouped: dict[int, list[PredictionResult]] = defaultdict(list)
    for result in population_rows:
        population_grouped[result.community_id].append(result)

    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    suffix = "整体" if mode == "overall" else re.sub(r"[^\w\u4e00-\u9fff-]", "_", next(iter(grouped.values()))[0].community.code)
    output_path = Path(project.path) / "output" / f"PreReport_run{run.id}_{suffix}_{REPORT_FORMAT_VERSION}_{timestamp}.docx"
    document = Document()
    _configure_document(document)
    document.core_properties.keywords = f"community-prediction-report;{REPORT_FORMAT_VERSION}"
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("完整社区设施需求预测与动态匹配结果报告")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("分社区报告" if mode == "community" else "整体报告").bold = True
    document.add_paragraph(f"工程名称：{project.name}")
    document.add_paragraph(f"生成时间：{datetime.now().strftime('%Y年%m月%d日 %H:%M:%S')}")
    first_rows = next(iter(grouped.values()))
    document.add_paragraph(f"预测时间范围：{first_rows[0].year}-{first_rows[-1].year}年")
    document.add_paragraph(f"数据说明：本报告由设施预测批次 #{run.id} 根据当前社区数据、关联人口预测结果和本批次设施参数自动生成，不使用固定社区或固定结果。")
    population_run = db.get(PredictionRun, population_run_id) if population_run_id else None
    if population_run and population_run.parameters_snapshot.get("migrationRateUnit") == "‰":
        document.add_paragraph("人口计算口径：出生率、死亡率和净迁入率均为千分率（‰）。出生人数=上年人口×出生率÷1000；死亡人数按上年分龄人口×分龄死亡率÷1000；净迁移人数按出生、死亡及年龄增长后的分龄人口×净迁入率÷1000。人数按半入规则取整，负净迁入率表示净迁出。")
    else:
        document.add_paragraph("历史人口计算口径：关联人口批次出生率、死亡率按千分率计算，净迁入率按当时的百分率计算；本报告保留历史预测结果，不使用新单位重新计算。")

    if mode == "overall":
        document.add_paragraph(f"本次报告包含 {len(grouped)} 个社区，社区清单如下：")
        community_table = document.add_table(rows=1, cols=4)
        community_table.style = "Table Grid"
        for cell, header in zip(community_table.rows[0].cells, ("序号", "社区编号", "社区名称", "参数参考地区")):
            _set_cell_text(cell, header, bold=True, color="FFFFFF")
            _shade(cell, BLUE)
        for index, rows in enumerate(grouped.values(), start=1):
            community = rows[0].community
            cells = community_table.add_row().cells
            for cell, value in zip(cells, (index, community.code, community.name, community.region)):
                _set_cell_text(cell, value)

    for index, rows in enumerate(grouped.values(), start=1):
        document.add_section(WD_SECTION.NEW_PAGE)
        community = rows[0].community
        population = population_grouped.get(community.id, [])
        if not population:
            raise ValueError(f"{community.code}-{community.name} 缺少人口预测结果，无法生成报告")
        _add_community_report(document, index, community, rows, population)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
