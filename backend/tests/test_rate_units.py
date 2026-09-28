from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.main import app
from app.database import Base
from app.models import Community, Project, PredictionRun
from app.defaults import AGE_GROUPS, MORTALITY_GROUPS, DEFAULT_POPULATION_PARAMETERS
from app.schemas import CommunityCreate
from app.services.rate_units import migrate_community_rate_units
from app.services.calculation_service import explain_result
from app.services.excel_service import export_community_template, parse_community, validate_community_workbook


def test_legacy_database_conversion_is_idempotent_and_preserves_history():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        p = Project(name="旧工程", path="test-only")
        db.add(p)
        db.flush()
        c = Community(project_id=p.id, code="OLD", name="旧社区", migration_rate_unit="%", migration_rates={"0-3岁":1.2,"18-59岁":-0.3,"80岁以上":0})
        new = Community(project_id=p.id, code="NEW", name="新社区", migration_rates={"0-3岁":8})
        run = PredictionRun(project_id=p.id, kind="population", status="success", parameters_snapshot={"unchanged":True})
        db.add_all([c,new,run]); db.commit()
        for _ in range(2):
            migrate_community_rate_units(db)
            assert c.migration_rates == {"0-3岁":12,"18-59岁":-3,"80岁以上":0}
            assert c.migration_rate_unit == "‰"
            assert new.migration_rates == {"0-3岁":8}
            assert run.parameters_snapshot == {"unchanged":True}
            assert c.facilities["_migrationUnitBackup"]["values"]["0-3岁"] == 1.2


def test_excel_new_roundtrip_legacy_conversion_and_conflict(tmp_path):
    path = tmp_path/"community.xlsx"
    payload = {"code":"RATE", "name":"单位社区", "age_counts":{"30":1000}, "migration_rates":dict.fromkeys(AGE_GROUPS,12)}
    export_community_template(path,2026,payload)
    assert parse_community(path)["migration_rates"] == payload["migration_rates"]
    wb=load_workbook(path)
    for row in range(8,16):
        assert "‰" in wb["社区基本信息"].cell(row,1).value
        wb["社区基本信息"].cell(row,1).value=wb["社区基本信息"].cell(row,1).value.replace("‰","%")
        wb["社区基本信息"].cell(row,2).value=1.2
    wb.save(path)
    assert validate_community_workbook(path)==[]
    data=parse_community(path)
    assert data["migration_rates"]==payload["migration_rates"]
    assert data["migration_rate_unit"]=="‰"
    wb["社区基本信息"].append(["社区0-3岁常住人口年均净迁入率（‰）",12])
    wb.save(path)
    assert validate_community_workbook(path)
    with pytest.raises(ValueError,match="冲突"):
        parse_community(path)


def test_excel_percentage_format_is_not_silently_scaled(tmp_path):
    path=tmp_path/'format.xlsx'
    export_community_template(path,2026,{'code':'FORMAT','name':'格式测试','age_counts':{'30':1000}})
    wb=load_workbook(path)
    wb['社区基本信息']['B8']=0.012
    wb['社区基本信息']['B8'].number_format='0.0%'
    wb.save(path)
    assert any('格式' in issue['field'] for issue in validate_community_workbook(path))
    with pytest.raises(ValueError,match='百分比格式'):
        parse_community(path)


@pytest.mark.parametrize("rate",[-1000,0,1000])
def test_migration_boundaries(rate):
    assert CommunityCreate(code="C",name="社区",migration_rates={"18-59岁":rate}).migration_rate_unit=="‰"


@pytest.mark.parametrize("rate",[-1001,1001,float('nan')])
def test_invalid_migration_boundaries(rate):
    with pytest.raises(ValueError):
        CommunityCreate(code="C",name="社区",migration_rates={"18-59岁":rate})


def test_three_rates_divide_by_1000_and_historical_explanation():
    with TestClient(app) as client:
        p=client.post('/api/projects',json={"name":"千分率算例-"+uuid4().hex}).json()
        root=f"/api/projects/{p['id']}"
        try:
            params=deepcopy(DEFAULT_POPULATION_PARAMETERS)
            params.update(startYear=2026,duration=1)
            params['birthRates']={'全国':{'2027':10}}
            params['deathRates']={'2027':dict.fromkeys(MORTALITY_GROUPS,0)}
            params['deathRates']['2027']['30-34岁']=5
            assert client.put(root+'/parameters/population',json={'data':params}).status_code==200
            ids=[]
            for code,rate in [('IN',20),('OUT',-20)]:
                c=client.post(root+'/communities',json={'code':code,'name':code,'region':'全国','population_base_year':2026,'age_counts':{'30':1000},'migration_rates':{'18-59岁':rate}})
                assert c.status_code==201,c.text
                ids.append(c.json()['id'])
            response=client.post(root+'/predictions/population',json={'seed':42})
            assert response.status_code==200,response.text
            run=response.json()
            for cid,total in zip(ids,[1025,985]):
                data=client.get(f"/api/runs/{run['id']}/calculation",params={'community_id':cid,'year':2027}).json()
                audit=data['audit']
                assert audit['births']==10
                assert sum(x['死亡人数'] for x in audit['deaths'])==5
                assert next(r['结果'] for r in data['rows'] if r['指标']=='预测总人口')==total
                assert '/ 1000)' in next(r['公式'] for r in data['rows'] if r['指标']=='净迁移人数')
                assert all('净迁入率（‰）' in row for row in audit['migration'])
                legacy=SimpleNamespace(kind='population',parameters_snapshot={},id=1,seed=42)
                historical=explain_result(legacy,SimpleNamespace(data={'calculationAudit':audit,'totalPopulation':total,'households':400},community_id=cid,year=2027))
                assert '/ 100)' in next(r['公式'] for r in historical['rows'] if r['指标']=='净迁移人数')
                assert '历史批次口径' in historical['notes'][1]
            bad=deepcopy(params);bad['birthRateUnit']='%'
            assert client.put(root+'/parameters/population',json={'data':bad}).status_code==422
        finally:
            client.delete(root)
