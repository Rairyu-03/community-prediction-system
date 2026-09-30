from copy import deepcopy
from io import BytesIO
from uuid import uuid4
import os
from pathlib import Path

WORKSPACE_ROOT = Path(__file__).resolve().parents[2]
os.environ.setdefault("DATABASE_URL", f"sqlite:///{(WORKSPACE_ROOT / 'backend' / 'data' / 'test_api.db').as_posix()}")
os.environ.setdefault("PROJECTS_ROOT", str(WORKSPACE_ROOT / "runtime" / "test_projects"))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.defaults import DEFAULT_FACILITY_PARAMETERS
from app.services.calculation_service import FACILITY_FORMULAS
from app.services.excel_service import parse_community, validate_community_workbook
from app.services.excel_service import export_community_template


def _complete_facility_parameters() -> dict:
    data = deepcopy(DEFAULT_FACILITY_PARAMETERS)
    data["meal"].update({"demandRate": 35, "diningAreaPerPerson": 1.2, "deliveryAreaPerPerson": 0.4, "kitchenAreaPerPerson": 0.35, "preparationAreaPerPerson": 0.25, "diningSpaceAreaPerPerson": 0.6})
    data["elderly"].update({"stationDemandRate": 20, "stationAreaPerPerson": 3, "rehabilitationRates": {"60-69岁": 18, "70-79岁": 25, "80岁以上": 32}, "entertainmentRates": {"60-69岁": 30, "70-79岁": 36, "80岁以上": 40}, "transportRates": {"60-69岁": 8, "70-79岁": 16, "80岁以上": 28}})
    data["childcare"].update({"nurseryDemandRate": 40, "nurseryAreaPerPerson": 8, "kindergartenDemandRate": 85, "kindergartenBuildingAreaPerPerson": 8.8, "kindergartenLandAreaPerPerson": 15, "outdoorAreaPerPerson": 4, "sharedAreaPerPerson": 2, "combinedRoomAreaPerPerson": 5.5, "activityRoomAreaPerPerson": 3.2, "bedroomAreaPerPerson": 2.8})
    data["parking"].update({"chargersPerHousehold": 0.25, "parkingSpacesPerHousehold": 1, "newEnergyVehiclesPerCapita": {str(year): 0.12 for year in range(2026, 2037)}, "vehiclesPerCapita": {str(year): 0.32 for year in range(2026, 2037)}})
    data["activity"].update({"publicAreaPerCapita": 0.3, "sportsAreaPerCapita": 0.1})
    return data


def test_summary_only_export_does_not_invent_residents(tmp_path):
    target = tmp_path/'summary.xlsx'
    export_community_template(target, 2026, {'code':'SUMMARY', 'name':'=1+1', 'population_base_year':2026, 'age_counts':{'0-3岁':5,'18-59岁':20}})
    wb=load_workbook(target,data_only=False)
    assert wb['社区人口列表'].max_row==1
    assert wb['社区基本信息']['B3'].data_type=='s'
    assert wb['社区基本信息']['B3'].value=='=1+1'
    assert validate_community_workbook(target)==[]
    parsed=parse_community(target)
    assert parsed['population_records']==[]
    assert parsed['age_counts']=={'0-3岁':5,'18-59岁':20}


def test_saved_community_export_roundtrip_and_scope(tmp_path):
    with TestClient(app) as client:
        project = client.post('/api/projects', json={'name': '导出测试-'+uuid4().hex}).json()
        root = f"/api/projects/{project['id']}"
        payload = {'code':'TEST-EXPORT', 'name':'自定义社区', 'region':'全国', 'households':0, 'average_household_size':2.5, 'population_base_year':2026, 'migration_rates':{'0-3岁':-1.5}, 'population_records':[{'序号':1,'性别':'女','出生年份':2020}, {'序号':2,'性别':'男','出生年份':1960}], 'facilities':{'公共活动场地':{'公共活动场地面积（平方米）':100, '多功能运动场地面积（平方米）':0}, '停车充电设施':{'停车位数量（个）':2,'机动车充电桩数量（个）':1}, '机动车充电桩列表':[{'机动车充电桩编号':'C-1','建设年份':2025,'经度':117.1,'纬度':39.1}]}, 'idle_spaces':[{'闲置空间编号':'SPACE-1','待分配建筑面积（平方米）':100,'待分配用地面积（平方米）':200,'经度':117.2,'纬度':39.2}]}
        saved = client.post(root+'/communities',json=payload)
        assert saved.status_code == 201, saved.text
        community=saved.json()
        content=client.get(root+'/communities/template',params={'community_id':community['id']}).content
        path=tmp_path/'export.xlsx'
        path.write_bytes(content)
        assert validate_community_workbook(path)==[]
        result=parse_community(path)
        for key in ('code','name','region','households','average_household_size','population_records','idle_spaces'):
            assert result[key]==community[key], key
        assert result['facilities']['机动车充电桩列表']==community['facilities']['机动车充电桩列表']
        assert result['facilities']['公共活动场地']==community['facilities']['公共活动场地']
        assert result['households']==0  # zero must not be replaced by inferred households
        blank=load_workbook(BytesIO(client.get(root+'/communities/template?blank=true').content))
        assert blank['社区基本信息']['B2'].value is None
        payload['code']='SECOND'
        client.post(root+'/communities',json=payload)
        assert client.get(root+'/communities/template').status_code==400
        assert client.get(root+'/communities/template?community_id=999999').status_code==404
        assert client.delete(root).status_code==204


def test_calculation_audit_is_frozen_and_reconciles():
    with TestClient(app) as client:
        p=client.post('/api/projects',json={'name':'公式测试-'+uuid4().hex}).json()
        root=f"/api/projects/{p['id']}"
        payload={'code':'FORMULA-1','name':'公式社区','region':'四川省','population_base_year':2026,'average_household_size':2.5,'age_counts':{'2':50,'5':80,'30':100,'65':120,'75':90,'85':30}, 'migration_rates':{'0-3岁':10, '60-69岁':-5}}
        c=client.post(root+'/communities',json=payload).json()
        pop=client.post(root+'/predictions/population',json={'seed':42})
        assert pop.status_code==200,pop.text
        run=pop.json()
        url=f"/api/runs/{run['id']}/calculation?community_id={c['id']}&year=2027"
        audit=client.get(url).json()
        assert audit['hasAudit']
        raw=audit['audit']
        values={r['指标']:r['结果'] for r in audit['rows']}
        assert values['预测总人口']==raw['previousPopulation']+raw['births']-sum(r['死亡人数'] for r in raw['deaths'])+sum(r['净迁移人数'] for r in raw['migration'])
        assert raw['birthRegion'] in ('全国','全国（中预测）')
        payload['average_household_size']=3
        client.put(root+f"/communities/{c['id']}",json=payload)
        assert client.get(url).json()==audit
        parameters=_complete_facility_parameters()
        assert client.put(root+'/parameters/facility',json={'data':parameters}).status_code==200
        fac=client.post(root+'/predictions/facility',json={'population_run_id':run['id']})
        assert fac.status_code==200,fac.text
        report=client.get(f"/api/runs/{fac.json()['id']}/calculation?community_id={c['id']}&year=2027").json()
        assert report['hasAudit']
        assert len(report['rows'])==len(FACILITY_FORMULAS)
        assert all('未登记' not in row['公式'] for row in report['rows'])
        assert report['audit']['populationRunId']==run['id']
        assert any('统一使用一个需求率' in note for note in report['notes'])
        assert any('不作为正式审批依据' in note for note in report['notes'])
        assert all('需求说明.docx' not in note and '第2.' not in note for note in report['notes'])
        assert client.get(f"/api/runs/{run['id']}/calculation?community_id={c['id']}&year=2099").status_code==404
        assert client.delete(root).status_code==204


def test_run_community_labels_follow_ids_not_code_order():
    from app.models import Community, PredictionRun, PredictionResult
    a = Community(id=8, code='Z', name='先创建的社区')
    b = Community(id=9, code='A', name='后创建的社区')
    run = PredictionRun(parameters_snapshot={'_communities':[{'id':9,'code':'A','name':b.name},{'id':8,'code':'Z','name':a.name}]})
    run.results = [PredictionResult(community_id=9, community=b), PredictionResult(community_id=8, community=a)]
    assert run.community_ids == [8,9]
    assert run.community_names == ['Z-先创建的社区','A-后创建的社区']
    run.parameters_snapshot = {}
    assert run.community_names == ['Z-先创建的社区','A-后创建的社区']


def test_pending_uniform_death_cannot_be_saved_and_new_horizon_runs():
    with TestClient(app) as client:
        project = client.post('/api/projects', json={'name':'年度联动-'+uuid4().hex}).json()
        root=f"/api/projects/{project['id']}"
        try:
            url=root+'/parameters/population'
            original=client.get(url).json()['data']
            pending={**original,'_pendingUniformDeath':{'0-4岁':0.5}}
            response=client.put(url,json={'data':pending})
            assert response.status_code==422
            assert '尚未应用' in response.text
            assert client.get(url).json()['data']==original
            data=deepcopy(original)
            data.update(startYear=2040,duration=3)
            reference=deepcopy(next(iter(data['deathRates'].values())))
            for year in range(2041,2044):
                data['deathRates'][str(year)]={**reference,'0-4岁':.5}
                for series in data['birthRates'].values():
                    series[str(year)]=4.5
            assert client.put(url,json={'data':data}).status_code==200
            c=client.post(root+'/communities',json={'code':'YEAR-2040','name':'远期模拟社区','region':'全国','population_base_year':2040,'age_counts':{'2':40,'30':200,'65':80},'average_household_size':2.5}).json()
            pop=client.post(root+'/predictions/population',json={'seed':17})
            assert pop.status_code==200,pop.text
            rows=client.get(f"/api/runs/{pop.json()['id']}/results").json()
            assert [r['year'] for r in rows]==[2041,2042,2043]
            for row in rows:
                assert row['data']['calculationAudit']['deaths'][0]['死亡率（‰）']==.5
            assert client.put(root+'/parameters/facility',json={'data':_complete_facility_parameters()}).status_code==200
            fac=client.post(root+'/predictions/facility',json={'population_run_id':pop.json()['id'],'generate_report':True,'report_mode':'community','community_id':c['id']})
            assert fac.status_code==200,fac.text
            assert fac.json()['forecast_years']==[2041,2042,2043]
            assert client.get(f"/api/runs/{fac.json()['id']}/report").status_code==200
        finally:
            client.delete(root)
