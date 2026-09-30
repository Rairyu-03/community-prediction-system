r"""Explicit live acceptance test: creates one NEW simulated project and retains it.

Run from workspace root: .venv\Scripts\python.exe -X utf8 scripts/test_two_community_permille_workflow.py
Never selects/deletes existing user projects. Only its newly-created runs are cleared.
"""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
import os
from pathlib import Path
import sys
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from app.defaults import AGE_GROUPS
from workflow_sample_data import SAMPLE_FACILITIES, SAMPLE_FACILITY_PARAMETERS, SAMPLE_POPULATION_PARAMETERS


def main():
    client = httpx.Client(base_url=os.environ.get("COMMUNITY_TEST_BASE_URL", "http://127.0.0.1:8000/api"), timeout=180)
    checks = []
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = ROOT / "runtime" / "acceptance" / stamp
    out.mkdir(parents=True)

    def call(method, path, expected=200, **kwargs):
        response = client.request(method, path, **kwargs)
        assert response.status_code == expected, (method, path, response.status_code, response.text[:700])
        return response

    def ok(message):
        checks.append(message)
        print(f"PASS {message}", flush=True)

    project = call("POST", "/projects", 201, json={"name": f"双社区千分率全流程验证工程（模拟）_{stamp}"}).json()
    prefix = f"/projects/{project['id']}"
    (out / "project.json").write_text(json.dumps(project, ensure_ascii=False, indent=2), encoding="utf-8")
    call("PATCH", prefix, json={"name": project["name"] + "_重命名验证"})
    call("PATCH", prefix, json={"name": project["name"]})
    call("POST", "/projects", 422, json={"name": "   "})
    ok("新建工程、重命名、空名称校验")

    population = deepcopy(SAMPLE_POPULATION_PARAMETERS)
    population.update(startYear=2026, duration=5)
    for kind, data in (("population", population), ("facility", SAMPLE_FACILITY_PARAMETERS)):
        call("PUT", prefix + f"/parameters/{kind}", json={"data": data})
        template = call("GET", prefix + f"/parameters/{kind}/template")
        call("POST", prefix + f"/parameters/{kind}/import", files={"file": (kind+".xlsx", template.content)})
    ok("人口/设施参数保存、下载、再导入")

    saved = []
    cases = [
        ("SIM-C", "青禾成长社区（模拟）", "天津市", [80,60,100,60,600,60,30,10], [.8,.6,.3,-.2,1.2,-.3,-.5,-.6]),
        ("SIM-A", "松龄养老社区（模拟）", "北京市", [20,20,40,30,350,260,200,80], [-.2,-.3,-.3,-.5,-1,.6,.8,.5]),
        
    ]
    for index, (code, name, region, counts, rates) in enumerate(cases):
        records = []
        for (lo, hi), count in zip(AGE_GROUPS.values(), counts):
            for k in range(count):
                age = lo + k % (min(hi, 95)-lo+1)
                records.append({"序号": len(records)+1, "性别": "男" if k%2 else "女", "出生年份": 2026-age, "数据来源": "自定义全流程测试模拟数据，非真实居民"})
        facilities = deepcopy(SAMPLE_FACILITIES)
        for values in facilities.values():
            if isinstance(values, dict):
                for key, value in values.items():
                    values[key] = round(value * (.2 + .1*index))
        facilities["公共活动场地"] = {"公共活动场地面积（平方米）": 120+index*80, "多功能运动场地面积（平方米）": 50+index*30}
        facilities["机动车充电桩列表"] = [{"机动车充电桩编号": f"{code}-C-{i+1}", "建设年份": 2020+i%6, "经度": 117.18+i*.00001, "纬度": 39.12+index*.001, "数据来源": "模拟"} for i in range(facilities["停车充电设施"]["机动车充电桩数量（个）"])]
        facilities["_dataNotes"] = ["此社区全部人口、设施、闲置空间、坐标及迁移率为测试模拟，不作为规划依据。设施预测参数同样为模拟值。"]
        payload = {"code":code,"name":name,"region":region,"households":round(sum(counts)/2.5),"average_household_size":2.5,"population_base_year":2026,"population_records":records,"migration_rates":dict(zip(AGE_GROUPS,[round(rate*10,4) for rate in rates])),"facilities":facilities,"idle_spaces":[{"闲置空间编号":f"{code}-S-{i}","待分配建筑面积（平方米）":100+i*150+index*30,"待分配用地面积（平方米）":200+i*200,"经度":117.18+i*.0001,"纬度":39.12+index*.001} for i in range(1,4)]}
        item = call("POST", prefix+"/communities",201,json=payload).json()
        saved.append(item)
        call("POST", prefix+"/communities",409,json=payload)
        assert sum(item["age_counts"].values()) == sum(counts)
    ok("创建两个完整社区：1000/1000人，设施、迁移率、充电桩、闲置空间全覆盖；重复编号拦截")
    assert call("GET",prefix+"/dashboard").json()["total_population"] == 2000
    assert len(call("GET",prefix+"/communities",params={"q":"SIM-A"}).json()) == 1
    assert len(call("GET",prefix+"/communities",params={"q":"不存在的社区"}).json()) == 0

    item = saved[0]
    item["population_records"].append({"序号":1001,"性别":"女","出生年份":2020})
    item["idle_spaces"].append({"闲置空间编号":"TEMP-SPACE","待分配建筑面积（平方米）":150,"待分配用地面积（平方米）":250,"经度":117.2,"纬度":39.1})
    updated = call("PUT",prefix+f"/communities/{item['id']}",json=item).json()
    assert sum(updated["age_counts"].values()) == 1001 and len(updated["idle_spaces"]) == 4
    item["population_records"].pop()
    item["idle_spaces"].pop()
    saved[0] = call("PUT",prefix+f"/communities/{item['id']}",json=item).json()
    invalid = deepcopy(item)
    invalid["population_records"][0]["出生年份"] = 2030
    call("PUT",prefix+f"/communities/{item['id']}",422,json=invalid)
    invalid = deepcopy(item)
    invalid["migration_rates"]["0-3岁"] = 1001
    call("PUT",prefix+f"/communities/{item['id']}",422,json=invalid)
    ok("搜索、人口/空间增删保存、非法出生年及迁移率校验")

    for item in saved:
        exported = call("GET",prefix+"/communities/template",params={"community_id":item["id"]}).content
        (out/f"{item['code']}.xlsx").write_bytes(exported)
        workbook = load_workbook(BytesIO(exported),data_only=True)
        assert workbook["社区人口列表"].max_row == len(item["population_records"])+1
        for row, rate in zip(range(8,16), item["migration_rates"].values()):
            assert workbook["社区基本信息"].cell(row,2).value == rate
            assert "‰" in workbook["社区基本信息"].cell(row,1).value
        imported = call("POST",prefix+"/communities/import",files={"file":(item["code"]+".xlsx",exported)}).json()
        assert imported["updated_count"] == 1 and imported["failed_count"] == 0, imported
    invalid_book = load_workbook(BytesIO(exported))
    invalid_book["社区基本信息"]["B8"] = "非法文本"
    stream = BytesIO()
    invalid_book.save(stream)
    rejected = call("POST",prefix+"/communities/import",files={"file":("invalid.xlsx",stream.getvalue())}).json()
    assert rejected["failed_count"] == 1 and rejected["issues"], rejected
    imported_book = load_workbook(BytesIO(exported))
    imported_book["社区基本信息"]["B2"] = "SIM-IMPORT-TEMP"
    imported_book["社区基本信息"]["B3"] = "导入新增删除验证（临时模拟）"
    stream = BytesIO()
    imported_book.save(stream)
    created = call("POST",prefix+"/communities/import",files={"file":("new.xlsx",stream.getvalue())}).json()
    assert created["created_count"] == 1 and created["failed_count"] == 0, created
    temporary = next(c for c in call("GET",prefix+"/communities").json() if c["code"] == "SIM-IMPORT-TEMP")
    call("DELETE",prefix+f"/communities/{temporary['id']}",204)
    assert len(call("GET",prefix+"/communities").json()) == 2
    call("GET",prefix+"/communities/template",400)
    ok("两个社区导出/导入更新一致、8项迁移率实值导出、异常字段报告、多社区未选中拦截")

    def predict():
        pop = call("POST",prefix+"/predictions/population",json={"seed":20260904}).json()
        assert pop["status"] == "success",pop
        # The streamlined UI no longer exposes a population-run selector. The
        # backend must therefore resolve the latest successful population run.
        facility = call("POST",prefix+"/predictions/facility",json={"seed":20260904,"generate_report":False}).json()
        assert facility["status"] == "success", facility
        generated = call("POST",f"/runs/{facility['id']}/report",json={"report_mode":"overall"}).json()
        assert generated["report_available"] is True
        return pop, facility

    pop, facility = predict()
    for run in (pop,facility):
        rows = call("GET",f"/runs/{run['id']}/results").json()
        assert len(rows) == 10 and set(r["year"] for r in rows) == set(range(2027,2032))
        by_id = {item["id"]:item for item in saved}
        assert run["community_names"] == [f"{by_id[i]['code']}-{by_id[i]['name']}" for i in run["community_ids"]],run
        for row in rows:
            if run["kind"] == "population":
                data = row["data"]
                audit = data["calculationAudit"]
                total = audit["previousPopulation"]+audit["births"]-sum(x["死亡人数"] for x in audit["deaths"])+sum(x["净迁移人数"] for x in audit["migration"])
                assert total == data["totalPopulation"]
                from app.services.prediction_service import round_half_up
                assert all(x["净迁移人数"] == round_half_up(x["基础存活人数"] * x["净迁入率（‰）"] / 1000) for x in audit["migration"])
            calculation = call("GET",f"/runs/{run['id']}/calculation",params={"community_id":row["communityId"],"year":row["year"]}).json()
            assert calculation
        result_file = call("GET",f"/runs/{run['id']}/download").content
        assert load_workbook(BytesIO(result_file),read_only=True).sheetnames
        (out/f"{run['kind']}-results.xlsx").write_bytes(result_file)
        backup = call("GET",f"/runs/{run['id']}/backup").content
        assert len(json.loads(ZipFile(BytesIO(backup)).read("results.json"))) == 10
        (out/f"{run['kind']}-backup.zip").write_bytes(backup)
    ok("双社区千分率5年人口/设施预测各10行；人口守恒、逐年计算依据、社区名称映射、结果和备份下载")

    def check_report(run, names):
        content = call("GET",f"/runs/{run['id']}/report").content
        archive = ZipFile(BytesIO(content))
        xml = ET.fromstring(archive.read("word/document.xml"))
        text = "".join(xml.itertext())
        assert all(name in text for name in names)
        assert "均为千分率" in text and "÷1000" in text
        assert all(str(year) in text for year in range(2027,2032))
        assert any(name.startswith("word/media/") for name in archive.namelist())
        (out/f"report-{run['id']}.docx").write_bytes(content)
        return text

    check_report(facility,[item["name"] for item in saved])
    call("DELETE",f"/runs/{pop['id']}",409)
    call("DELETE",prefix+f"/communities/{saved[0]['id']}",409)
    call("POST",prefix+"/runs/clear",409,json=[])
    call("DELETE",f"/runs/{facility['id']}",204)
    call("DELETE",f"/runs/{pop['id']}",204)
    pop, facility = predict()
    ids = [r["id"] for r in call("GET",prefix+"/runs").json()]
    cleared = call("POST",prefix+"/runs/clear",json=ids).json()
    assert cleared["deleted_count"] == len(ids)
    assert call("GET",prefix+"/runs").json() == []
    assert len(call("GET",prefix+"/communities").json()) == 2
    ok("整体报告含两个社区/五年/图表；依赖删除拦截、单条删除、一键清除及社区数据保留")

    pop, facility = predict()
    check_report(facility,[item["name"] for item in saved])
    for item in saved:
        single = call("POST",prefix+"/predictions/facility",json={"generate_report":False}).json()
        call("POST",f"/runs/{single['id']}/report",json={"report_mode":"community","community_id":item["id"]})
        text = check_report(single,[item["name"]])
        assert all(other["name"] not in text for other in saved if other["id"] != item["id"])
    ok("重新生成并保留人口批次、整体报告和两个分社区报告，报告无跨社区串入")
    summary = {"project":project,"communities":[{"id":c["id"],"code":c["code"],"name":c["name"],"population":sum(c["age_counts"].values())} for c in saved],"checks":checks,"runs":call("GET",prefix+"/runs").json()}
    (out/"summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
    print("RESULT "+str(out),flush=True)


if __name__ == "__main__":
    main()
