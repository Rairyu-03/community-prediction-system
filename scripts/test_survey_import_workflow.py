r"""Import the supplied comprehensive survey workbook and retain a real acceptance project.

Run from the workspace root:
    .venv\Scripts\python.exe -X utf8 scripts\test_survey_import_workflow.py
"""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import sys
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from fastapi.testclient import TestClient
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.demo_data import DEMO_FACILITY_PARAMETERS  # noqa: E402
from app.main import app  # noqa: E402


SOURCE = ROOT / "社区数据" / "610103-D 《社区基本信息与服务设施信息调研表》.xlsx"


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    evidence = ROOT / "runtime" / "acceptance" / f"survey_import_{stamp}"
    evidence.mkdir(parents=True, exist_ok=True)
    checks: list[str] = []

    def ok(message: str) -> None:
        checks.append(message)
        print(f"PASS {message}", flush=True)

    with TestClient(app) as client:
        project_response = client.post("/api/projects", json={"name": f"建科大社区调研表全流程验证_{stamp}"})
        assert project_response.status_code == 201, project_response.text
        project = project_response.json()
        base = f"/api/projects/{project['id']}"

        with SOURCE.open("rb") as source:
            response = client.post(
                base + "/communities/import",
                files={"file": (SOURCE.name, source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert response.status_code == 200, response.text
        import_report = response.json()
        assert import_report["success_count"] == 1 and import_report["created_count"] == 1, import_report
        community = import_report["community"]
        assert community["code"] == "610103-D"
        assert community["name"] == "文艺路街道建科大社区"
        assert community["region"] == "陕西省"
        assert sum(community["age_counts"].values()) == 6550
        assert community["households"] == 2974
        assert community["facilities"]["托育设施"]["现有托儿所托育人数（人）"] == 54
        assert community["facilities"]["托育设施"]["现有幼儿园托育人数（人）"] == 348
        assert len(community["facilities"]["机动车充电桩列表"]) == 3
        assert len(import_report["warnings"]) >= 4
        ok("非标准综合调研表自动识别、转换、校验并入库")

        exported = client.get(base + "/communities/template", params={"community_id": community["id"]})
        assert exported.status_code == 200, exported.text
        exported_book = load_workbook(BytesIO(exported.content), data_only=True, read_only=True)
        assert {"社区基本信息", "社区人口列表", "托育设施", "机动车充电桩列表"} <= set(exported_book.sheetnames)
        assert exported_book["社区人口列表"].max_row == 6551
        exported_book.close()
        (evidence / "converted-community-data.xlsx").write_bytes(exported.content)
        roundtrip = client.post(
            base + "/communities/import",
            files={"file": ("converted-community-data.xlsx", exported.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert roundtrip.status_code == 200 and roundtrip.json()["updated_count"] == 1, roundtrip.text
        ok("转换后的系统社区现状数据可下载并无损回导")

        population = client.get(base + "/parameters/population").json()["data"]
        population.update({"startYear": 2026, "duration": 5})
        response = client.put(base + "/parameters/population", json={"data": population})
        assert response.status_code == 200, response.text

        facility = deepcopy(DEMO_FACILITY_PARAMETERS)
        facility["isDemo"] = False
        facility["parameterNature"] = "流程验证模拟标准；正式项目应替换为当地现行配置标准"
        for reference in facility.get("references", {}).values():
            if isinstance(reference, dict):
                reference["source"] = "流程验证模拟标准"
        response = client.put(base + "/parameters/facility", json={"data": facility})
        assert response.status_code == 200, response.text
        assert client.get(base + "/quality", params={"kind": "population"}).json()["ready"] is True
        ok("人口参数通过预测前质量检查，模拟设施标准已保存")

        population_run_response = client.post(base + "/predictions/population", json={"seed": 20260914})
        assert population_run_response.status_code == 200, population_run_response.text
        population_run = population_run_response.json()
        population_rows = client.get(f"/api/runs/{population_run['id']}/results").json()
        assert len(population_rows) == 5
        assert {row["year"] for row in population_rows} == {2027, 2028, 2029, 2030, 2031}
        assert all("totalPopulation" in row["data"] and "households" in row["data"] for row in population_rows)
        assert all(len(row["data"]["ageGroups"]) == 8 for row in population_rows)
        calculation = client.get(
            f"/api/runs/{population_run['id']}/calculation",
            params={"community_id": community["id"], "year": 2027},
        )
        assert calculation.status_code == 200 and calculation.json(), calculation.text
        ok("完成5年人口预测，包含8个年龄段、总人数、总户数及计算说明")

        facility_quality = client.get(base + "/quality", params={"kind": "facility"})
        assert facility_quality.status_code == 200 and facility_quality.json()["ready"] is True, facility_quality.text
        ok("人口预测完成后，设施预测前质量检查通过")

        facility_run_response = client.post(
            base + "/predictions/facility",
            json={
                "population_run_id": population_run["id"],
                "seed": 20260914,
                "generate_report": True,
                "report_mode": "overall",
            },
        )
        assert facility_run_response.status_code == 200, facility_run_response.text
        facility_run = facility_run_response.json()
        facility_rows = client.get(f"/api/runs/{facility_run['id']}/results").json()
        assert len(facility_rows) == 5
        assert all(row["data"].get("metrics") for row in facility_rows)
        assert any(row["data"].get("siteSuggestions") for row in facility_rows)
        ok("完成5年设施预测并生成基于缺口与现状的选址建议")

        for run in (population_run, facility_run):
            download = client.get(f"/api/runs/{run['id']}/download")
            assert download.status_code == 200 and download.content.startswith(b"PK")
            (evidence / f"{run['kind']}-results.xlsx").write_bytes(download.content)

        report_response = client.get(f"/api/runs/{facility_run['id']}/report")
        assert report_response.status_code == 200 and report_response.content.startswith(b"PK")
        report_archive = ZipFile(BytesIO(report_response.content))
        report_text = "".join(ET.fromstring(report_archive.read("word/document.xml")).itertext())
        assert community["name"] in report_text
        assert all(str(year) in report_text for year in range(2027, 2032))
        assert any(name.startswith("word/media/") for name in report_archive.namelist())
        assert "设施需求预测结果" in report_text and "选址" in report_text
        (evidence / "facility-prediction-report.docx").write_bytes(report_response.content)
        ok("报告包含社区名称、完整5年结果及人口/设施图表")

        dashboard = client.get(base + "/dashboard").json()
        assert dashboard["community_count"] == 1 and dashboard["total_population"] == 6550
        assert len(client.get(base + "/runs").json()) == 2
        summary = {
            "project": project,
            "community": {"id": community["id"], "code": community["code"], "name": community["name"]},
            "importWarnings": import_report["warnings"],
            "populationRunId": population_run["id"],
            "facilityRunId": facility_run["id"],
            "facilityParameterNote": facility["parameterNature"],
            "checks": checks,
        }
        (evidence / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"PROJECT_ID {project['id']}")
        print(f"EVIDENCE {evidence}")


if __name__ == "__main__":
    main()
