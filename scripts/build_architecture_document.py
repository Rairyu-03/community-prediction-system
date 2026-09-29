from pathlib import Path
from math import atan2, cos, sin, pi
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / 'runtime' / 'architecture-document'
WORK.mkdir(parents=True, exist_ok=True)
OUT = ROOT / '完整社区预测系统项目架构设计说明.docx'
NAVY, TEAL, INK, MUTED = '#173E59', '#058B99', '#183548', '#506A7C'
PALE, BORDER = '#EFF7FA', '#CBDDE6'
FONT = 'C:/Windows/Fonts/msyh.ttc'
BOLD = 'C:/Windows/Fonts/msyhbd.ttc'

class Diagram:
    def __init__(self, name, h):
        self.name=name
        self.im=Image.new('RGB',(1800,h),'white')
        self.d=ImageDraw.Draw(self.im)
    def text(self,x,y,s,size=30,bold=False,color=INK,anchor='mm'):
        self.d.text((x,y),s,font=ImageFont.truetype(BOLD if bold else FONT,size),fill=color,anchor=anchor)
    def box(self,x,y,w,h,title,lines=(),fill=PALE,accent=TEAL):
        self.d.rounded_rectangle((x,y,x+w,y+h),radius=18,fill=fill,outline=BORDER,width=2)
        self.d.rounded_rectangle((x+16,y+20,x+23,y+h-20),radius=3,fill=accent)
        self.text(x+w/2,y+40,title,33,True)
        for i,s in enumerate(lines): self.text(x+w/2,y+86+i*39,s,27,color=MUTED)
    def arrow(self,points,color=TEAL,width=5):
        self.d.line(points,fill=color,width=width,joint='curve')
        x,y=points[-1]; px,py=points[-2]; a=atan2(y-py,x-px)
        self.d.polygon([(x,y),(x-17*cos(a-.45),y-17*sin(a-.45)),(x-17*cos(a+.45),y-17*sin(a+.45))],fill=color)
    def save(self):
        path=WORK/(self.name+'.png');self.im.save(path,dpi=(300,300));return path

g=Diagram('整体架构图',1120)
g.box(50,20,810,140,'独立使用',['浏览器工作台  /  源码部署'])
g.box(940,20,810,140,'嵌入大系统',['React 组件  /  iframe  /  宿主事件'])
g.arrow([(450,160),(450,205)]);g.arrow([(1350,160),(1350,205)])
g.box(50,205,1700,195,'表现层   React 18 + TypeScript + Vite',[
    '工程入口与管理  ·  社区数据与参数表单  ·  预测结果与计算说明',
    '自定义 SVG 组合折线图  ·  社区与批次切换  ·  浅色与深色主题'])
g.arrow([(900,400),(900,458)]);g.text(1135,430,'HTTP / JSON / 文件传输',25)
g.box(50,458,1700,135,'接口层   FastAPI + Pydantic',['参数校验  ·  API 路由  ·  可选令牌校验  ·  CORS'])
for x in (325,900,1475): g.arrow([(x,593),(x,640)])
g.box(50,640,550,190,'业务服务', ['工程与社区管理','导入校验与预测前检查'])
g.box(625,640,550,190,'计算服务',['人口逐年演进与设施测算','面积匹配与计算依据'])
g.box(1200,640,550,190,'成果与迁移服务',['Excel 与 Word 报告','批次备份与工程迁移'])
for x in (325,900,1475):g.d.line([(x,830),(x,865)],fill=TEAL,width=5)
g.d.line([(325,865),(1475,865)],fill=TEAL,width=5)
g.arrow([(450,865),(450,900)]);g.arrow([(1350,865),(1350,900)])
g.box(50,900,810,180,'SQLAlchemy → SQLite',['工程 / 参数 / 社区 / 批次 / 结果','结构化数据与计算快照'])
g.box(940,900,810,180,'工程文件目录',['导入 Excel / 结果 Excel / Word / ZIP','数据库保存文件路径与业务关联'])
architecture=g.save()

g=Diagram('部署架构图',610)
g.box(35,30,550,170,'开发环境',['浏览器 → Vite  5173','HTTP API → FastAPI  8000'])
g.box(625,30,550,170,'生产部署',['静态页面 → 反向代理','FastAPI 服务 → 持久化目录'])
g.box(1215,30,550,170,'大系统集成',['宿主加载组件或 iframe','调用同一套后端 API'])
g.arrow([(900,200),(900,285)])
g.box(220,285,1360,170,'业务数据持久化',['DATABASE_URL + PROJECTS_ROOT','SQLite 数据库 + projects 工程目录 + trash 回收备份'])
g.text(900,535,'源码开发、生产部署和嵌入集成复用同一业务服务',31,True)
deploy=g.save()

g=Diagram('业务流程图',790)
stages=[('01  选择工程','内置案例只读查看','新建或打开普通工程'),('02  录入社区','表单填写 / Excel 导入','检查人口与设施现状'),('03  人口预测','配置年份与人口参数','质量检查后逐年计算'),('04  设施预测','选择人口批次与标准','计算五类需求及缺口')]
for i,(t,a,b) in enumerate(stages):
    x=35+i*450;g.box(x,40,380,205,t,[a,b])
    if i<3:g.arrow([(x+380,140),(x+435,140)])
g.arrow([(1575,245),(1575,340),(900,340),(900,395)])
g.box(465,395,870,185,'05  查看成果与导出',['按社区、类型、批次查看组合趋势图','查看计算逻辑 → 下载 Excel 与整体或分社区报告'])
g.arrow([(900,580),(900,635)])
g.box(465,635,870,125,'06  备份或调整参数重新预测',['新批次保留历史，跨设备使用工程包迁移'])
workflow=g.save()

g=Diagram('选址分析图',420)
for i,(t,ls) in enumerate([
    ('计算新增缺口',['预测需求减现有供给','用于选址的负缺口取零']),
    ('逐一匹配空间',['比较建筑和用地面积','统计达标数与覆盖率']),
    ('确定候选方案',['达标时优先较小余量','未达标时选最高覆盖率']),
    ('生成文字建议',['写入年份、缺口与空间','附现场复核事项'])]):
    x=30+450*i;g.box(x,60,380,230,t,ls)
    if i<3:g.arrow([(x+380,175),(x+435,175)])
g.text(900,365,'当前实现为面积条件筛选，尚未执行 GIS 服务半径或跨设施空间分配',28,color=MUTED)
site=g.save()

g=Diagram('数据实体关系图',850)
g.box(595,20,610,165,'projects  工程',['主键 id  ·  工程名称  ·  文件路径'])
g.box(25,300,535,170,'parameter_sets  参数',['project_id → projects.id','kind + data JSON'])
g.box(635,300,535,170,'communities  社区',['project_id → projects.id','人口 / 设施 / 闲置空间 JSON'])
g.box(1245,300,535,170,'prediction_runs  批次',['project_id → projects.id','类型 / 状态 / 参数快照 / 文件路径'])
g.arrow([(740,185),(740,240),(290,240),(290,300)])
g.arrow([(900,185),(900,300)]);g.arrow([(1060,185),(1060,240),(1510,240),(1510,300)])
for x in (285,965,1510):g.text(x,268,'1 对多',24,color=TEAL)
g.box(625,635,700,170,'prediction_results  年度结果',['community_id + run_id + year','data JSON 保存指标、计算依据和选址分析'])
g.arrow([(900,470),(900,635)]);g.text(800,555,'1 对多',24,color=TEAL)
g.arrow([(1510,470),(1510,720),(1325,720)]);g.text(1530,585,'1 对多',24,color=TEAL)
g.text(287,590,'工程内社区编号唯一',27,color=MUTED)
g.text(287,635,'同工程同类参数唯一',27,color=MUTED)
g.text(287,680,'同批次同社区同年结果唯一',27,color=MUTED)
erd=g.save()

g=Diagram('报告与备份数据流图',570)
g.box(35,25,540,190,'预测批次和年度结果',['参数快照 + 计算审计','按工程与社区定位'])
g.box(630,25,540,190,'成果生成',['Pillow 绘制人口图表','python-docx 排版 Word 报告'])
g.box(1225,25,540,190,'文件落盘与下载',['openpyxl 写入 Excel','Word 图表嵌入报告文件'])
g.arrow([(575,120),(630,120)]);g.arrow([(1170,120),(1225,120)])
g.box(35,335,540,190,'导出工程 ZIP',['清单 + 结构化数据','input / output 工程文件'])
g.box(630,335,540,190,'导入时重新映射',['重建工程、社区和批次 ID','恢复日期与目标文件路径'])
g.box(1225,335,540,190,'恢复为可访问副本',['历史结果继续查看','成果文件继续下载'])
g.arrow([(575,435),(630,435)]);g.arrow([(1170,435),(1225,435)])
g.text(900,272,'跨设备恢复使用备份包，源码与业务数据应分别管理',28,True)
artifacts=g.save()

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.67);sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
for name,size in [('Normal',11),('Title',24),('Subtitle',12),('Heading 1',18),('Heading 2',12.5),('Caption',9)]:
    s=doc.styles[name];s.font.name='微软雅黑';s.font.size=Pt(size)
    s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.line_spacing=1.1
    if name.startswith('Heading'):s.font.bold=True
doc.styles['Normal'].paragraph_format.widow_control=True
doc.styles['Subtitle'].font.italic=False
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('项目架构设计  ·  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(70,70,70)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.core_properties.title='完整社区设施需求预测与动态匹配系统项目架构设计说明'
doc.core_properties.subject='技术架构 数据架构 功能设计 部署与使用'
doc.core_properties.author='项目组'

def p(text='',bold=False):
    q=doc.add_paragraph();r=q.add_run(text);r.bold=bold;return q
def h(text,level=1):doc.add_heading(text,level)
def page(title):
    q=doc.add_heading(title,1);q.paragraph_format.page_break_before=True
def fig(path,caption,width=7.08):
    q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(3)
    q.add_run().add_picture(str(path),width=Inches(width));q.paragraph_format.keep_with_next=True
    shape=doc.inline_shapes[-1];shape._inline.docPr.set('descr',caption)
    q=doc.add_paragraph(caption,'Caption');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173E59' if ri==0 else ('F0F6F9' if ri%2 else 'FFFFFF'));tc.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tc.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'70' if side in ('top','bottom') else '95');e.set(qn('w:type'),'dxa');margins.append(e)
            tc.append(margins)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(1);q.paragraph_format.line_spacing=1.1
                for r in q.runs:
                    r.font.size=Pt(9.5);r.bold=ri==0
                    r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '183548')
    p().paragraph_format.space_after=Pt(1)
    return t

doc.add_paragraph('完整社区设施需求预测\n与动态匹配系统','Title')
doc.add_paragraph('项目架构设计说明','Subtitle')
p('版本 1.0    编制日期 2026年9月7日    面向项目维护与系统集成',False)
p('本系统以工程组织社区数据、人口参数、设施标准和预测成果，采用 React 前端与 Python 后端分离设计。它既可独立运行，也可嵌入大系统；人口及设施结果由参数化计算生成，并形成可追溯的图表、Excel 和 Word 报告。')
fig(architecture,'图 1  系统整体分层架构')
p('核心设计：页面负责录入和展示，后端负责校验与计算，SQLite 保存结构化业务数据，工程目录保存成果文件。当前为单体后端分层架构，业务服务在同一 Python 应用内运行。')

page('1 技术选型与部署形态')
p('技术选择围绕前后端复用、可维护部署与大系统集成展开。开发环境前后端独立启动，生产环境由反向代理统一提供静态页面与 API。')
table(['层次','当前技术','职责与选择理由'],[
('界面','React 18 + TypeScript','组件化表单与结果页面；类型约束减少接口字段不一致。'),
('构建与样式','Vite + CSS','生成独立网页和嵌入组件；CSS 变量实现浅色、深色主题。'),
('图表','自定义 SVG','组合折线、悬停数值、图例显隐及线性或对数尺度切换。'),
('接口与服务','Python + FastAPI + Uvicorn','HTTP API、文件上传下载和同步预测服务。'),
('校验与配置','Pydantic + pydantic-settings','输入模型、字段范围校验、环境变量配置。'),
('持久化','SQLAlchemy 2 + SQLite','ORM 管理五类核心实体；SQLite 无需额外安装。'),
('成果生成','openpyxl / python-docx / Pillow','读写 Excel、生成 Word，并将人口图表嵌入报告。'),
('部署','Vite 构建 + Uvicorn + 反向代理','前端静态资源与后端 API 分层部署，便于监控、扩容和集成。'),
],[.8,2.0,4.25])
fig(deploy,'图 2  开发与交付部署形态')
p('PostgreSQL 为可选配置路径，通过 DATABASE_URL 和对应驱动切换。当前开发与单实例部署使用 SQLite；多人部署仍需单独验证并发、权限和备份策略。')

page('2 功能模块与实现分工')
p('系统按“工程 → 社区 → 预测批次 → 年度结果”组织业务。一个普通工程可录入多个社区，各社区共享该工程的人口及设施参数，并保存各自的现状与净迁入率。')
table(['功能入口','用户可完成的操作','主要实现'],[
('工程入口与管理','新建、打开、切换、重命名、删除工程；导入导出工程；系统备份恢复。','App.tsx\nproject_service / portability_service'),
('社区数据输入','手动新建或导入 Excel；查看导入成功、更新、失败和异常字段；下载选中社区现状数据。','CommunityEditor\nexcel_service / schemas'),
('社区管理','搜索编号、名称或地区；查看编辑基础、设施、人口和闲置空间信息；进入社区预测结果。','CommunityEditor\ncommunity_service'),
('人口预测参数','基准年和时长；各地区出生率；共用或逐年死亡率；参数保存、导入和导出。','ParameterForms\nforecastParameters'),
('设施预测参数','五类设施的需求率、面积配置标准、车辆年度参数及来源；保存和 Excel 往返。','ParameterForms\nvalidation_service'),
('预测结果','按社区、类型、批次查看组合图；设施按类别及单位切换；查看计算逻辑及说明。','TrendLineChart / CalculationPanel\ncalculation_service'),
('报告与历史管理','下载结果及整体或分社区报告；备份、删除和清空历史；检查人口批次引用关系。','report_service / run_service\nprediction_service'),
],[1.15,3.65,2.25])
h('组件与服务的边界',2)
p('CommunityApi 统一封装请求、文件上传和错误信息。App.tsx 管理工程与页面状态，独立组件负责社区编辑、参数表单和计算说明。后端 main.py 暴露 API，services 目录承载具体业务，models.py 定义数据关系。')
p('内置“默认预测项目”用于庆有西里社区只读演示。用户录入真实数据和调整参数时使用普通工程；演示参数、模拟数据和正式项目数据应在业务层明确区分。')

page('3 业务流程与人口预测设计')
fig(workflow,'图 3  从工程准备到预测成果的完整业务流程')
h('人口年度演进过程',2)
p('1  按社区人口数据建立 0 至 120 岁年龄分布。已有居民出生年份按基准年转换为年龄；120 岁作为开放末龄桶。')
p('2  读取预测年度出生率，用上一年总人口乘出生率并除以 1000，得到出生人数。地区出生率优先采用社区参考地区，缺少该地区序列时尝试全国序列。')
p('3  按上一年各死亡年龄组人数、死亡率计算死亡人数，再分配到组内年龄；存活人口递增一岁，并加入新生人口。')
p('4  按年龄组存活人数乘净迁入率并除以 1000，计算净迁移人数；正值迁入，负值迁出。汇总人口及户数后进入下一年。')
p('5  保存每社区每年结果、出生和死亡明细、净迁移明细、使用参数及算法版本。随机分配使用种子和社区编号派生的随机序列，便于相同输入复现。')
p('单位与年份：出生率、死亡率、净迁入率均为 ‰；基准年 2026、预测时长 5 年对应 2027—2031 年。人口取整采用半入取整，预测户数由预测人口除以户均人数获得。',True)

page('4 设施测算与选址建议设计')
p('设施预测以已成功的人口预测批次为输入，逐社区、逐年执行五类测算。地方设施标准由用户维护，参考来源随指标写入结果；设施服务需求率仍使用百分率，与人口参数的千分率分开处理。')
table(['设施类别','输入与计算要点','输出示例'],[
('助餐设施','老年人口 × 助餐需求率，再乘堂食、送餐及分功能面积标准。','助餐人数；食堂、厨房、配餐、用餐面积；现状缺口。'),
('养老服务','按老年年龄组加权计算失能、康复、娱乐、接送需求；配置服务站容量和面积。','服务人数、失能占比、站点容量及面积缺口。'),
('托育服务','0—3 岁与 4—6 岁人口分别乘需求率及相应面积标准。','托儿所、幼儿园建筑、用地、室外及功能空间面积。'),
('停车充电','户均配置法与人均车辆保有量法；充电桩按建设年和寿命扣除到期数量。','充电桩与停车位缺口；两种口径结果并列展示。'),
('公共活动','总人口乘公共活动和运动场地人均标准。','公共活动面积、多功能运动场地面积。'),
],[1.05,3.45,2.55])
fig(site,'图 4  闲置空间匹配与建议生成过程')
h('选址建议的实际范围',2)
p('系统对助餐、养老服务站、公共活动和多功能运动场地进行面积匹配。建议中的年份、需求缺口、候选空间编号、达标数量和覆盖率由计算生成，复核事项使用设施类别规则文案。')
p('当前坐标用于标识候选位置，尚未计算路网距离、GIS 服务半径、消防或规划适宜性，也未实现跨设施争用同一空间的联合分配。托育和停车类别已有需求预测，但没有对应的完整空间选址模型。')
p('年度车辆参数优先使用精确年份；中间缺年采用线性插值，区间外采用最近端点。指标缺口可以为负数表示供给富余，新增选址需求按不小于零处理。')

page('5 数据库与文件存储设计')
fig(erd,'图 5  核心数据库实体关系')
table(['实体或存储','保存内容','关键约束或访问方式'],[
('projects','工程名称、路径、创建与更新时间。','工程名称和路径唯一。'),
('parameter_sets','人口、设施参数及参数来源路径。','project_id 与 kind 组合唯一。'),
('communities','基础信息、分龄人数、居民列表、净迁入率、设施和闲置空间。','工程内 code 唯一；复杂结构采用 JSON。'),
('prediction_runs','预测种类、状态、种子、参数快照、错误、时间及结果文件路径。','设施批次在快照中记录所用人口批次 ID。'),
('prediction_results','每社区每年指标、计算明细及选址分析。','run_id、community_id、year 组合唯一。'),
('工程文件目录','input 下的导入文件；output 下的结果 Excel、报告及图表资源。','文件内容保存在磁盘；通过 API 下载。'),
],[1.5,3.25,2.3])
p('图表的数据来自年度结果 JSON，网页 SVG 按需绘制。Word 报告中的人口图表由 Pillow 生成并嵌入文档；报告和 Excel 二进制内容没有作为数据库 BLOB 保存。因此完整迁移需要同时包含数据库记录与工程文件。')

page('6 结果追溯与备份恢复设计')
fig(artifacts,'图 6  成果生成与工程迁移的数据流')
h('预测成果如何随社区变化',2)
p('人口和设施计算逐个读取当前工程内的社区，按所选预测年度生成结果。报告使用指定设施批次及关联的人口结果，按整体或选中社区过滤，插入人口图表、逐年设施表、参考来源和选址建议。报告沿用统一版式，数据和分析内容随批次变化。')
p('“查看计算逻辑及说明”从该年度结果的计算审计数据和批次参数快照读取公式、代入值与结果。后续修改社区或参数后，需要重新执行预测生成新批次；历史结果继续保留其原有计算依据。')
table(['操作','包含内容','恢复或保护机制'],[
('批次备份','批次清单、参数快照、结果 JSON、Excel 和报告。','用于归档与核查；当前未提供独立批次 ZIP 导入入口。'),
('工程导出导入','工程、参数、社区、批次、年度结果及工程文件。','目标端创建副本并重映射 ID、日期和文件路径。'),
('系统备份恢复','所有当前工程的工程包及系统清单。','恢复为多个新副本；不会覆盖同名工程。'),
('删除与清空','按批次删除或清除本工程历史。','先保留内部备份；被设施引用的人口批次不可直接删除。'),
],[1.1,3.0,2.95])
p('恢复边界：系统备份是当前工程的逻辑备份，不是整台设备镜像，不含程序环境配置和所有回收目录。当前恢复按工程依次提交，尚未实现整个系统包的全局原子回滚。')

page('7 操作流程与部署使用')
table(['步骤','操作位置','输入与完成标准'],[
('1  进入工程','入口页面或顶部新建工程','输入名称创建普通工程；只体验功能时打开默认案例。'),
('2  录入数据','社区数据输入','导入社区 Excel 或手动填写；完成基础、人口、设施、充电桩及闲置空间信息。'),
('3  核对参数','人口预测参数','确认人口基准年与预测范围；填齐所需地区出生率、死亡率。人口三类率均为 ‰。'),
('4  人口预测','人口预测参数底部','保存后执行，处理检查错误；成功后可按社区查看年龄组合趋势。'),
('5  设施预测','设施预测参数底部','填写本地设施标准及来源，选择人口批次，指定整体或分社区报告并执行。'),
('6  查看成果','预测结果','选择社区、预测类型和历史批次；设施选择类别及单位；查看计算说明并下载成果。'),
('7  跨机迁移','工程管理','导出工程或下载系统备份；在另一设备导入或恢复。'),
],[1.05,1.6,4.4])
h('源码运行与生产部署',2)
p('开发者可使用根目录“一键启动.bat”或 scripts/start-dev.ps1 启动，前端访问 5173 端口，后端 API 使用 8000 端口。首次运行前需要安装 Python 与 Node.js 依赖。')
p('生产环境先执行前端构建，再由 Web 服务器提供静态资源；FastAPI 由 Uvicorn 或进程管理器运行，并通过反向代理暴露 /api。业务数据目录应使用持久化磁盘并定期备份。')
p('默认开发数据位于 backend/data/community_prediction.db 和 runtime/projects。可通过 DATABASE_URL 与 PROJECTS_ROOT 指定生产数据位置；跨环境迁移优先使用工程导出或系统备份。')

page('8 接口与嵌入集成设计')
p('模块提供独立页面、iframe 和 React 组件三种接入形式，复用同一套 HTTP API。React 模式通过 CommunityPredictionModule 或 mountCommunityPredictionModule 挂载，并可传入工程、接口地址和主题。')
table(['接入项','可配置内容','作用'],[
('组件与通信','mode、projectId、apiBase、defaultTheme；hostOrigin、onEvent、hostContext','控制布局、工程和主题；通过回调、CustomEvent 或 postMessage 通知宿主。'),
('认证与跨域','token / HOST_API_TOKEN；CORS_ORIGINS','可选共享令牌及允许访问来源；未配置令牌时不启用令牌校验。'),
('存储配置','DATABASE_URL、PROJECTS_ROOT','配置数据库与工程文件位置，并指向受备份保护的持久化存储。'),
],[1.05,3.0,3.0])
h('主要接口分组',2)
table(['API 路径示例','用途'],[
('GET 或 POST /api/projects','列出工程或创建工程。'),
('/api/projects/{id}/communities','社区查询、新建；具体社区路径支持编辑和删除。'),
('/api/projects/{id}/parameters/{kind}','读取和保存人口或设施参数。'),
('GET /api/projects/{id}/quality','预测前质量检查，kind 指定 population 或 facility。'),
('POST /api/projects/{id}/predictions/{kind}','同步执行人口或设施预测，返回批次信息。'),
('/api/runs/{id}/…','results 读取结果；calculation 读取计算说明；download 或 report 下载成果。'),
('/api/projects/{id}/export\n/api/projects/import-archive','导出或导入完整工程包。'),
('/api/system/backup 或 restore','全部工程备份和按副本恢复。'),
],[3.45,3.6])
p('集成注意：当前令牌是共享接口令牌，尚无用户账户、角色权限或工程级授权体系。hostContext 属于上下文传递，不自动形成权限边界。宿主应配置可信来源，并结合自身身份体系完成权限接入。')

page('9 维护边界与扩展方向')
h('当前架构的适用范围',2)
p('系统适合单机运行、社区案例测算和模块集成验证。计算逻辑由 Python 规则和用户参数驱动，不依赖大模型。更换普通工程和社区后，可沿用数据录入、人口预测、设施测算和报告生成流程；地方标准仍需要用户提供和确认。')
p('当前预测在 API 请求内同步执行，没有后台任务队列与进度恢复机制；大批量社区、长年限或多人同时预测需进一步评估。SQLite 用于本机数据管理，工程 ID 关联提供业务隔离，但不等同于多租户安全隔离。')
h('关键设计约束',2)
table(['设计点','当前实现与维护要求'],[
('参数质量','页面预测入口先调用质量检查；服务端另有输入及计算校验。API 集成应主动调用 quality，不能假定所有预检规则都在预测接口统一拦截。'),
('历史追溯','已有批次保留计算结果、参数与关键审计数据；当前快照不等同于完整居民身份与全量录入文件版本库。'),
('数据库演进','启动时执行建表和已有字段的增量迁移。后续复杂结构变更宜引入正式版本迁移管理。'),
('备份可靠性','工程数据与文件共同迁移。系统级一致性快照、恢复全局回滚和超大备份流式处理仍可继续完善。'),
('选址模型','面积匹配已实现；服务半径、地理可达性、规划约束和空间联合分配属于后续扩展。'),
],[1.25,5.8])
page('10 源码组织与验证记录')
p('以下定位用于维护者快速找到接口、算法、页面和发布脚本。扩展业务时应同时检查参数校验、计算过程、结果说明和报告输出，避免只修改其中一处。')
h('源码定位',2)
table(['文件或目录','维护职责'],[
('frontend/src/App.tsx\nParameterForms.tsx / CommunityEditor.tsx','页面编排、结果图表、参数表单与社区资料。'),
('frontend/src/api.ts / embed.tsx / hostBridge.ts','API 请求、组件挂载和宿主事件。'),
('backend/app/main.py / schemas.py / models.py','接口路由、请求校验及数据库模型。'),
('backend/app/services/prediction_service.py\ncalculation_service.py / report_service.py','人口设施计算、计算依据和报告生成。'),
('quality_service.py / portability_service.py\nexcel_service.py / run_service.py','预测预检、备份迁移、Excel 和历史记录管理。'),
('scripts/start-dev.ps1 / stop-dev.ps1\nfrontend/vite*.ts','开发启动、前端构建与嵌入组件配置。'),
],[3.7,3.35])
table(['既有验证记录','验证范围'],[
('后端自动化测试','2026年9月6日记录为 56 项通过，覆盖 API、计算、导入导出、单位与备份迁移等。'),
('前端构建','TypeScript 检查、独立页面生产构建和嵌入组件构建通过。'),
('双社区流程','青禾成长与松龄养老两个模拟社区，各 1000 人；2027—2031 年人口、设施结果及整体与分社区报告。'),
('源码运行验证','前后端启动、健康接口、默认案例、页面访问及质量检查通过。'),
],[1.5,5.55])
p('本设计对应 2026年9月7日的项目代码。此前交付验证记录包含 56 项后端测试、前端与嵌入构建及双社区全流程验收；这些记录说明已验证的功能范围，不替代后续版本回归或并发性能测试。')

for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.save(OUT)
print(OUT)
