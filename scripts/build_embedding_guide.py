from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "交付文档"
ASSET_DIR = ROOT / "runtime" / "embedding-guide"
OUT_PATH = OUT_DIR / "完整社区预测模块嵌入大系统接入说明.docx"
BLUE = "155F78"
TEAL = "078B91"
LIGHT_BLUE = "EAF5F7"
PALE_BLUE = "F4F9FB"
GRAY = "D9D9D9"
TEXT = "173B4D"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf" if bold else r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "#4B8297", width: int = 5) -> None:
    draw.line((*start, *end), fill=color, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    spread = 0.55
    points = [
        end,
        (int(end[0] - length * math.cos(angle - spread)), int(end[1] - length * math.sin(angle - spread))),
        (int(end[0] - length * math.cos(angle + spread)), int(end[1] - length * math.sin(angle + spread))),
    ]
    draw.polygon(points, fill=color)


def rounded_box(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], title: str, detail: str, fill: str, outline: str = "#79A9B8") -> None:
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    draw.text((x1 + 28, y1 + 25), title, fill="#103D52", font=font(28, True))
    lines = detail.split("\n")
    for index, line in enumerate(lines):
        draw.text((x1 + 28, y1 + 75 + index * 34), line, fill="#4C6F7F", font=font(21))


def make_architecture_diagram(path: Path) -> None:
    image = Image.new("RGB", (1800, 940), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 42), "大系统嵌入架构", fill="#111111", font=font(42, True))
    draw.text((70, 102), "推荐由统一网关提供同源地址，模块通过 iframe 或 React 组件加载", fill="#5E7682", font=font(23))

    rounded_box(draw, (70, 205, 405, 430), "宿主大系统", "菜单按钮与权限入口\n用户身份与工程上下文\n统一导航和页面框架", "#E8F4F7")
    rounded_box(draw, (530, 205, 900, 430), "社区预测前端模块", "六功能单页面工作台\n人口与设施柱状图\nExcel 与 Word 文件交互", "#E5F7F3", "#39A69D")
    rounded_box(draw, (1030, 205, 1390, 430), "FastAPI 服务", "业务接口与输入校验\n预测计算与报告生成\nBearer Token 可选校验", "#EDF2FB", "#6B8FC2")
    rounded_box(draw, (1460, 205, 1730, 430), "数据层", "SQLite 数据库\n工程文件目录\n结果与备份文件", "#F5F1FB", "#9D83BE")
    arrow(draw, (405, 315), (530, 315))
    arrow(draw, (900, 315), (1030, 315))
    arrow(draw, (1390, 315), (1460, 315))
    draw.text((425, 275), "iframe 或组件", fill="#39697B", font=font(19, True))
    draw.text((920, 275), "HTTPS JSON API", fill="#39697B", font=font(19, True))

    rounded_box(draw, (300, 590, 740, 800), "宿主接收模块事件", "module ready 与错误状态\n工程创建切换与删除\n预测完成失败与主题变化", "#F8FAFC")
    rounded_box(draw, (990, 590, 1430, 800), "统一反向代理", "同源路径 /community prediction\nAPI 路径 /community api\nTLS CORS 与访问控制", "#F8FAFC")
    arrow(draw, (715, 590), (715, 485), "#0A8F91")
    arrow(draw, (715, 485), (405, 485), "#0A8F91")
    arrow(draw, (405, 485), (405, 430), "#0A8F91")
    arrow(draw, (1210, 590), (1210, 430), "#5479A6")
    draw.text((458, 505), "postMessage 或 onEvent", fill="#0A7679", font=font(20, True))
    draw.text((1230, 500), "推荐生产部署", fill="#466A91", font=font(20, True))
    image.save(path)


def make_sequence_diagram(path: Path) -> None:
    image = Image.new("RGB", (1800, 980), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 38), "嵌入页面访问时序", fill="#111111", font=font(42, True))
    actors = [(150, "用户"), (520, "宿主系统"), (930, "预测模块"), (1370, "API 与数据层")]
    for x, label in actors:
        draw.rounded_rectangle((x - 100, 120, x + 100, 190), radius=18, fill="#EAF5F7", outline="#7DA9B7", width=3)
        bbox = draw.textbbox((0, 0), label, font=font(24, True))
        draw.text((x - (bbox[2] - bbox[0]) / 2, 139), label, fill="#17475B", font=font(24, True))
        draw.line((x, 190, x, 900), fill="#B8CCD4", width=2)
    steps = [
        (245, 150, 520, "点击社区预测按钮"),
        (325, 520, 930, "加载模块并传入 projectId"),
        (405, 930, 1370, "携带 Token 请求模块配置和工程数据"),
        (485, 1370, 930, "返回社区 参数 批次和结果"),
        (565, 930, 520, "发送 module ready 事件"),
        (645, 930, 1370, "用户执行预测或下载报告"),
        (725, 1370, 930, "返回预测批次及输出文件"),
        (805, 930, 520, "发送 prediction completed 事件"),
    ]
    for y, start_x, end_x, label in steps:
        arrow(draw, (start_x, y), (end_x, y), "#168A93" if end_x > start_x else "#557C93", 4)
        text_x = min(start_x, end_x) + 18
        draw.text((text_x, y - 34), label, fill="#34596A", font=font(19, True))
    image.save(path)


def set_font(run, name: str = "Microsoft YaHei", size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "5")
        item.set(qn("w:color"), GRAY)
        node.append(item)
    tbl_pr.append(node)


def cell_text(cell, text: str, *, bold=False, color=TEXT, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(str(text))
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Cm(widths[i])
        shade(cell, BLUE)
        cell_text(cell, header, bold=True, color="FFFFFF", size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[i])
            if row_index % 2:
                shade(cell, PALE_BLUE)
            cell_text(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.4
    if bold_lead:
        lead = p.add_run(bold_lead)
        set_font(lead, size=10.5, bold=True, color="000000")
    run = p.add_run(text)
    set_font(run, size=10.5, color=TEXT)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    set_font(p.add_run(text), size=10.2, color=TEXT)


def add_code(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.05
    for index, line in enumerate(code.strip().splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, name="Consolas", size=8.3, color="253F4C")
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F6F7")
    p_pr.append(shd)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    architecture = ASSET_DIR / "嵌入架构图.png"
    sequence = ASSET_DIR / "访问时序图.png"
    make_architecture_diagram(architecture)
    make_sequence_diagram(sequence)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for style_name, size, before, after in (("Title", 27, 0, 14), ("Heading 1", 18, 18, 10), ("Heading 2", 14, 14, 7), ("Heading 3", 11.5, 10, 5)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_style_properties = styles["Title"].element.get_or_add_pPr()
    title_border = title_style_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_style_properties.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title_properties = title._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run("完整社区预测模块嵌入大系统接入说明"), size=27, bold=True, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    set_font(subtitle.add_run("单页面工作台接入架构 部署配置 接口参数与联调验收"), size=12, color="526E7C")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(meta.add_run("版本 1.0    2026年9月17日"), size=10, color="6B7F89")
    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    add_body(doc, "本文面向负责将完整社区预测模块接入现有业务平台的产品、前端、后端、运维和测试人员。当前模块已经统一为六功能单页面工作台，不包含旧入口页和左侧导航。大系统只需提供一个菜单按钮，并采用 iframe 或 React 组件方式加载模块，即可使用社区数据、人口预测、设施预测、结果报告、社区管理和工程管理全部功能。")
    add_body(doc, "推荐方案是通过统一反向代理把宿主系统、预测前端和预测 API 发布为同源路径，再使用 iframe 嵌入。该方案与宿主技术栈无关，隔离性和升级稳定性最好。宿主为 React 且希望深度统一样式时，可改用组件接入。", bold_lead="接入结论  ")
    doc.add_page_break()

    doc.add_heading("1 模块范围和接入边界", level=1)
    add_table(doc, ["层级", "本模块负责", "宿主系统负责"], [
        ["入口", "提供完整单页面工作台", "提供菜单按钮 路由 权限入口和返回路径"],
        ["业务", "社区数据录入 参数配置 人口预测 设施预测 图表 报告和工程管理", "决定哪些用户可访问 可选传入指定工程"],
        ["数据", "通过 FastAPI 读写数据库和工程文件", "提供部署目录 备份策略和必要的用户上下文"],
        ["安全", "支持 Bearer Token 和 CORS 白名单", "统一登录 HTTPS 网关鉴权和来源校验"],
        ["联动", "通过 postMessage 或 onEvent 发出状态事件", "监听事件并更新菜单 状态或审计记录"],
    ], [2.4, 7.1, 7.1])
    add_body(doc, "模块不是一个只展示图表的静态页面。预测计算、Excel 导入导出、Word 报告、工程备份和 SQLite 数据都依赖后端服务，因此不能只复制一个 HTML 文件完成生产接入。")

    doc.add_heading("2 整体接入架构", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(architecture), width=Inches(6.8))
    caption = doc.add_paragraph("图 1  大系统嵌入架构")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(caption.runs[0], size=9, color="526E7C")
    add_body(doc, "宿主系统只控制模块入口和上下文，不应直接读写社区预测数据库。所有业务读写都通过模块 API 完成；这可以保持预测规则、数据校验、报告生成和备份逻辑的一致性。")

    doc.add_heading("3 接入方式选择", level=1)
    add_table(doc, ["方式", "适用场景", "优点", "注意事项"], [
        ["iframe 推荐", "宿主技术栈不限 希望快速稳定接入", "隔离性好 发布独立 升级风险低", "配置同源代理或 CORS 监听 postMessage"],
        ["React 组件", "宿主已经使用 React 18 且需要统一页面生命周期", "可直接传入回调和上下文 无 iframe 边界", "宿主需提供 React 依赖 并同步加载模块 CSS"],
        ["新窗口跳转", "只需增加入口按钮 暂不要求页面内嵌", "实施最快 对宿主改动最少", "用户体验不如页内嵌入 宿主与模块状态联动较弱"],
    ], [2.5, 5.3, 5.1, 5.1])
    doc.add_heading("4 iframe 接入步骤", level=1)
    doc.add_heading("4.1 推荐地址规划", level=2)
    add_table(doc, ["用途", "示例地址", "说明"], [
        ["宿主系统", "https://portal.example.com", "大系统现有访问地址"],
        ["预测模块页面", "https://portal.example.com/community-prediction/", "反向代理到预测前端"],
        ["预测 API", "https://portal.example.com/community-api", "反向代理到 FastAPI 的 /api"],
    ], [3.0, 7.0, 6.7])
    doc.add_page_break()
    doc.add_heading("4.2 宿主按钮和页面", level=2)
    add_code(doc, """
<button id="open-community-prediction">社区预测</button>
<section id="community-prediction-panel" hidden>
  <iframe id="community-prediction-frame"
    title="完整社区预测"
    allow="clipboard-write"
    sandbox="allow-scripts allow-same-origin allow-forms allow-downloads">
  </iframe>
</section>
""")
    add_code(doc, """
const moduleOrigin = "https://portal.example.com";
const frame = document.querySelector("#community-prediction-frame");

document.querySelector("#open-community-prediction").onclick = () => {
  const params = new URLSearchParams({
    projectId: "9",
    apiBase: "/community-api",
    hostOrigin: window.location.origin
  });
  frame.src = `/community-prediction/?${params}`;
  document.querySelector("#community-prediction-panel").hidden = false;
};

window.addEventListener("message", (event) => {
  if (event.origin !== moduleOrigin) return;
  if (event.data?.source !== "community-prediction") return;
  console.log("社区预测模块事件", event.data.type, event.data.payload);
});
""")
    doc.add_heading("4.3 页面参数", level=2)
    add_table(doc, ["参数", "是否必填", "作用", "示例"], [
        ["projectId", "否", "进入指定预测工程 不传则进入可用的非演示工程", "9"],
        ["apiBase", "同源默认否", "指定 API 根地址 默认使用 /api", "/community-api"],
        ["token", "启用后必填", "作为 Bearer Token 访问 API 下载地址也会使用该令牌", "临时或网关签发令牌"],
        ["hostOrigin", "推荐", "限制模块 postMessage 的目标来源", "https://portal.example.com"],
    ], [2.4, 2.5, 7.0, 4.8])
    add_body(doc, "生产环境优先采用同源反向代理，并由网关完成登录校验。若必须把 token 放入 iframe 地址，应使用 HTTPS、缩短有效期并避免写入访问日志；当前后端也支持通过 HOST_API_TOKEN 配置共享访问令牌。")

    doc.add_heading("5 React 组件接入步骤", level=1)
    add_body(doc, "组件产物位于 frontend/dist-embed。运行 npm run build:embed 后会生成 community-prediction-module.js 和 community-prediction-module.css。构建时 React ReactDOM 和 ReactDOM Client 被设为外部依赖，由宿主系统提供。")
    add_code(doc, """
import { mountCommunityPredictionModule } from "./community-prediction-module.js";
import "./community-prediction-module.css";

const unmount = mountCommunityPredictionModule(
  document.getElementById("community-prediction-root"),
  {
    mode: "embedded",
    apiBase: "/community-api",
    projectId: 9,
    token: accessToken,
    hostOrigin: window.location.origin,
    hostContext: { userId: currentUser.id, organizationId: currentOrg.id },
    defaultTheme: "light",
    onEvent(event) {
      console.log(event.type, event.payload);
    }
  }
);

// 宿主页面销毁时调用
unmount();
""")
    add_table(doc, ["组件属性", "类型", "说明"], [
        ["mode", "standalone 或 embedded", "嵌入大系统时使用 embedded"],
        ["apiBase", "string", "FastAPI 根地址 默认 /api"],
        ["token", "string", "可选访问令牌 自动写入 Authorization 请求头"],
        ["projectId", "number", "可选 初始工程编号"],
        ["hostContext", "object", "随 API 请求写入 X Host Context 请求头"],
        ["hostOrigin", "string", "postMessage 目标来源"],
        ["defaultTheme", "light 或 dark", "模块初始主题"],
        ["onEvent", "function", "组件模式的宿主事件回调"],
    ], [3.0, 4.4, 9.4])
    doc.add_heading("6 模块事件", level=1)
    add_table(doc, ["事件名", "触发时机", "典型用途"], [
        ["module-ready", "模块和初始工程数据加载完成", "关闭宿主加载动画 记录当前工程"],
        ["module-error", "启动或服务连接失败", "显示统一错误提示或上报告警"],
        ["project-created", "创建工程完成", "同步宿主业务对象与工程编号"],
        ["project-changed", "用户切换工程", "更新地址栏或宿主面包屑"],
        ["project-renamed", "工程重命名", "刷新宿主显示名称"],
        ["project-deleted", "工程删除", "清理宿主关联记录"],
        ["prediction-completed", "人口或设施预测完成", "刷新状态 徽标或通知"],
        ["prediction-failed", "预测执行失败", "提示用户并记录失败原因"],
        ["prediction-run-deleted", "预测批次删除", "更新宿主历史状态"],
        ["theme-changed", "用户切换明暗主题", "可选 同步宿主主题偏好"],
    ], [4.0, 6.0, 6.8])
    add_body(doc, "iframe 事件结构为 source 等于 community-prediction、version 等于 1，并包含 type 和 payload。宿主必须同时校验 event.origin 和 source，不能只根据事件名执行敏感操作。")

    doc.add_heading("7 后端部署与安全配置", level=1)
    add_table(doc, ["配置项", "用途", "建议"], [
        ["CORS_ORIGINS", "允许跨域访问模块 API 的宿主来源", "生产环境只填写明确的 HTTPS 来源 多个来源用逗号分隔"],
        ["HOST_API_TOKEN", "启用 API Bearer Token 校验", "生产环境设置高强度值并由网关或宿主安全传递"],
        ["DATABASE_URL", "数据库连接地址", "当前默认 SQLite 单实例部署可直接使用"],
        ["PROJECTS_ROOT", "工程文件 结果 报告和备份目录", "使用持久化磁盘并纳入备份"],
    ], [3.2, 6.0, 7.6])
    add_bullet(doc, "生产环境必须使用 HTTPS，并由反向代理限制请求体大小 超时时间和下载路径。")
    add_bullet(doc, "生产部署应由 Uvicorn 或进程管理器运行后端，并通过反向代理统一提供静态页面与 /api。")
    add_bullet(doc, "SQLite 适合当前单实例业务。如果未来出现多节点部署或大量并发写入，应迁移到 PostgreSQL，并把工程文件迁移到共享存储。")
    add_bullet(doc, "不要让宿主前端直接访问数据库文件。工程迁移应使用系统备份或工程 ZIP。")

    doc.add_heading("8 推荐反向代理规则", level=1)
    add_code(doc, """
location /community-prediction/ {
    alias /srv/community-prediction/frontend/dist/;
    try_files $uri $uri/ /community-prediction/index.html;
}

location /community-api/ {
    proxy_pass http://127.0.0.1:8000/api/;
    proxy_set_header Host $host;
    proxy_set_header X-Forwarded-Proto $scheme;
    proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;
    proxy_read_timeout 180s;
}
""")
    add_body(doc, "如果宿主地址与预测模块地址完全同源，通常不需要浏览器跨域配置。若分属不同域名，必须同时配置 CORS_ORIGINS、HTTPS 和严格的 postMessage 来源校验。")

    doc.add_heading("9 访问时序", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(sequence), width=Inches(6.8))
    caption = doc.add_paragraph("图 2  嵌入页面访问与事件回传时序")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(caption.runs[0], size=9, color="526E7C")

    doc.add_heading("10 联调验收清单", level=1)
    for item in [
        "宿主按钮可打开单页面工作台 页面不出现旧入口页和左侧导航",
        "projectId 有效时进入指定工程 无效时返回明确错误或使用默认工程策略",
        "社区 Excel 可导入和下载 参数可保存 人口与设施预测均可执行",
        "人口和设施结果显示年度分组柱状图 图例 悬停和尺度切换可用",
        "结果 Excel Word 报告 工程 ZIP 和系统备份均可下载",
        "module-ready prediction-completed 和 module-error 能被宿主正确接收",
        "跨域 Token 下载和刷新页面后仍能正常访问",
        "生产环境数据目录有持久化和备份 恢复演练通过",
    ]:
        add_bullet(doc, "□ " + item)

    doc.add_heading("11 常见问题", level=1)
    add_table(doc, ["现象", "原因", "处理方法"], [
        ["页面可以打开但数据显示连接失败", "apiBase 错误 API 未启动或代理路径未保留 /api", "检查 /api/health 和浏览器网络请求"],
        ["浏览器提示跨域错误", "宿主来源未加入 CORS_ORIGINS", "增加精确来源后重启后端"],
        ["下载文件返回 401", "下载链接未携带有效 token", "确认传入 token 与 HOST_API_TOKEN 一致 或改为同源网关代理"],
        ["宿主收不到事件", "hostOrigin 错误 iframe 不是父子关系或来源校验不匹配", "使用宿主真实 origin 并检查 message 事件"],
        ["本机可以访问但其他设备无法访问", "仍在使用 127.0.0.1 或网关未开放服务", "使用正式域名和反向代理，并检查防火墙与监听地址"],
        ["组件页面无样式", "只加载了 JS 没加载 CSS", "同时引入 community-prediction-module.css"],
    ], [4.2, 5.4, 7.4])

    doc.add_page_break()
    doc.add_heading("附录 A 源码下载与一键启动", level=1)
    add_body(doc, "协作成员从 GitHub 获取源码后，可以通过根目录的一键启动脚本完成本地运行。该方式直接基于源码，用于开发、联调和功能验证。")
    add_table(doc, ["步骤", "操作", "完成标准"], [
        ["1 准备环境", "安装 Python 3.11 或更高版本、Node.js 20 或更高版本，并加入 PATH", "命令行可执行 python 或 py，以及 node 和 corepack"],
        ["2 获取源码", "git clone 仓库，或下载 ZIP 后完整解压", "根目录中可见 backend、frontend、scripts 和一键启动.bat"],
        ["3 一键启动", "双击根目录的一键启动.bat", "脚本自动创建 .venv、安装依赖、启动服务并打开浏览器"],
        ["4 访问系统", "浏览器打开 http://127.0.0.1:5173", "前端页面和 http://127.0.0.1:8000/api/health 均可访问"],
        ["5 停止服务", "双击根目录的停止服务.bat", "5173 与 8000 端口释放"],
    ], [2.4, 8.0, 6.4])
    add_bullet(doc, "首次启动需要联网下载 Python 和前端依赖，后续启动仍会校验依赖以适配代码更新。")
    add_bullet(doc, "不要从压缩包预览窗口直接运行，也不要把 .venv、node_modules 或运行数据库提交到 GitHub。")
    add_bullet(doc, "若提示端口占用，先运行停止服务.bat，再关闭占用 5173 或 8000 端口的其他程序。")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("完整社区预测模块嵌入大系统接入说明"), size=8.5, color="748993")
    doc.core_properties.title = "完整社区预测模块嵌入大系统接入说明"
    doc.core_properties.subject = "单页面工作台 iframe 和 React 组件接入说明"
    doc.core_properties.author = "完整社区预测项目组"
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build())
